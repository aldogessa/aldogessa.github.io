"""
==========================================================================
AGis - Aldo Gessa
SCRIPT UNICO – QGIS 3.44.15 (versione GeoPackage)
Intersezione particelle × layer tematici + Compilazione CDU
Senza layer fisici, con template Word esistente.
Versione Unificata – Ottobre 2026
==========================================================================
"""

# ============================================================
# IMPORT
# ============================================================

import os
import sys
import subprocess
import platform
import re
import processing

from qgis.core import (
    QgsProject,
    QgsField,
    QgsFeature,
    QgsWkbTypes,
    Qgis
)

from qgis.utils import iface

from qgis.PyQt.QtCore import (
    Qt,
    QVariant,
    QSettings,
    QUrl
)

from qgis.PyQt.QtGui import (
    QIcon,
    QPixmap
)

from qgis.PyQt.QtNetwork import (
    QNetworkAccessManager,
    QNetworkRequest
)

from qgis.PyQt.QtWidgets import (
    QDialog,
    QVBoxLayout,
    QHBoxLayout,
    QTreeWidget,
    QTreeWidgetItem,
    QTextEdit,
    QPushButton,
    QLineEdit,
    QLabel,
    QMessageBox,
    QProgressBar,
    QApplication,
    QComboBox,
    QFileDialog,
    QFrame
)


# ============================================================
# CONTROLLO / INSTALLAZIONE PYTHON-DOCX
# ============================================================

def ensure_python_docx():

    try:

        import docx

        print(
            f"python-docx già installato: {docx.__version__}"
        )

        return True

    except ImportError:

        pass


    # --------------------------------------------------------
    # Ricava il vero Python utilizzato da QGIS
    # --------------------------------------------------------

    qgis_root = os.path.dirname(
        os.path.dirname(
            sys.executable
        )
    )

    python_version = (
        f"Python{sys.version_info.major}{sys.version_info.minor}"
    )

    python_exe = os.path.join(
        qgis_root,
        "apps",
        python_version,
        "python.exe"
    )


    print(
        "python-docx non trovato."
    )

    print(
        "Python QGIS:"
    )

    print(
        python_exe
    )


    if not os.path.isfile(
        python_exe
    ):

        QMessageBox.critical(
            None,
            "Dipendenza mancante",
            "Impossibile individuare il Python utilizzato da QGIS."
        )

        return False


    # --------------------------------------------------------
    # Installazione automatica
    # --------------------------------------------------------

    iface.messageBar().pushMessage(
        "CDU",
        "Installazione del componente python-docx in corso... Attendere.",
        level=Qgis.Info,
        duration=0
    )


    # Forza QGIS a visualizzare immediatamente il messaggio
    # prima di avviare pip, che è un processo bloccante.

    QApplication.processEvents()


    try:

        result = subprocess.run(
            [
                python_exe,
                "-m",
                "pip",
                "install",
                "python-docx"
            ],
            capture_output=True,
            text=True,
            creationflags=subprocess.CREATE_NO_WINDOW
        )


    except Exception as e:

        QMessageBox.critical(
            None,
            "Installazione fallita",
            f"Impossibile avviare l'installazione di python-docx.\n\n{e}"
        )

        return False


    # Rimuove il messaggio "installazione in corso"

    iface.messageBar().clearWidgets()


    # --------------------------------------------------------
    # Verifica risultato pip
    # --------------------------------------------------------

    print(
        "=== INSTALLAZIONE PYTHON-DOCX ==="
    )

    print(
        result.stdout
    )


    if result.stderr:

        print(
            result.stderr
        )


    print(
        "Return code:",
        result.returncode
    )


    if result.returncode != 0:

        QMessageBox.critical(
            None,
            "Installazione fallita",
            "L'installazione di python-docx non è riuscita.\n\n"
            "Controlla la console Python di QGIS per i dettagli."
        )

        return False


    # --------------------------------------------------------
    # Verifica finale
    # --------------------------------------------------------

    try:

        import docx

        print(
            "python-docx installato correttamente."
        )

        print(
            "Versione:",
            docx.__version__
        )

        print(
            "Percorso:",
            docx.__file__
        )

        return True


    except ImportError as e:

        QMessageBox.critical(
            None,
            "Errore",
            "python-docx è stato installato ma QGIS non riesce a importarlo "
            "RIAVVIA QGIS PER TERMINARE L'INSTALLAZIONE DELL'AZIONE CDU.\n\n"
            f"{e}"
        )

        return False


# ============================================================
# DIALOGO DI SELEZIONE LAYER
# ============================================================

class LayerSelectionDialog(QDialog):

    def __init__(
        self,
        polygon_layers,
        parcel_layers,
        output_tables,
        active_layer,
        parent=None
    ):

        super().__init__(
            parent
        )

        self.active_layer = active_layer

        self.setWindowTitle(
            "Intersezioni CDU – Selezione layer"
        )

        self.resize(
            800,
            700
        )


        # ====================================================
        # CONTENITORE PRINCIPALE BIANCO
        # ====================================================

        self.container = QFrame()

        self.container.setSizePolicy(
            self.container.sizePolicy().Expanding,
            self.container.sizePolicy().Expanding
        )

        self.container.setStyleSheet(
            "background-color: white;"
        )


        self.settings = QSettings(
            "AGis",
            "IntersezioniCDU"
        )


        self.polygon_layers = polygon_layers
        self.parcel_layers = parcel_layers
        self.output_tables = output_tables


        # ====================================================
        # LAYOUT ESTERNO DELLA FINESTRA
        # ====================================================

        outer_layout = QVBoxLayout()

        outer_layout.setContentsMargins(
            0,
            0,
            0,
            0
        )

        self.setLayout(
            outer_layout
        )


        # ====================================================
        # CONTENITORE LOGO
        # ====================================================

        logo_frame = QFrame()

        logo_frame.setMinimumHeight(
            80
        )

        logo_frame.setStyleSheet(
            "background-color: white;"
        )

        logo_frame.setSizePolicy(
            logo_frame.sizePolicy().Expanding,
            logo_frame.sizePolicy().Fixed
        )


        # Layout interno del logo

        logo_layout = QHBoxLayout(
            logo_frame
        )

        logo_layout.setContentsMargins(
            0,
            0,
            0,
            0
        )

        logo_layout.setAlignment(
            Qt.AlignLeft | Qt.AlignVCenter
        )


        # ====================================================
        # LOGO
        # ====================================================

        self.logo_label = QLabel()

        self.logo_label.setAlignment(
            Qt.AlignLeft | Qt.AlignVCenter
        )

        self.logo_label.setMinimumHeight(
            80
        )

        logo_layout.addWidget(
            self.logo_label
        )


        # ====================================================
        # DOWNLOAD LOGO AGIS
        # ====================================================

        logo_url = (
            "https://raw.githubusercontent.com/"
            "aldogessa/aldogessa.github.io/main/"
            "docs/risorse/immagini/LogoAgisModulo.png"
        )


        self.network_manager = QNetworkAccessManager(
            self
        )


        request = QNetworkRequest(
            QUrl(logo_url)
        )


        reply = self.network_manager.get(
            request
        )


        def logo_download_finished():

            if reply.error() == reply.NetworkError.NoError:

                pixmap = QPixmap()


                if pixmap.loadFromData(
                    reply.readAll()
                ):

                    self.logo_label.setPixmap(
                        pixmap.scaled(
                            260,
                            75,
                            Qt.KeepAspectRatio,
                            Qt.SmoothTransformation
                        )
                    )

                else:

                    self.logo_label.setText(
                        "AGis – Aldo Gessa"
                    )


            else:

                self.logo_label.setText(
                    "AGis – Aldo Gessa"
                )


            reply.deleteLater()


        reply.finished.connect(
            logo_download_finished
        )


        # ====================================================
        # CONTENITORE PRINCIPALE CDU
        # ====================================================

        self.container = QFrame()

        self.container.setSizePolicy(
            self.container.sizePolicy().Expanding,
            self.container.sizePolicy().Expanding
        )

        self.container.setStyleSheet(
            "background-color: white;"
        )


        # Layout delle due colonne CDU

        main_layout = QHBoxLayout(
            self.container
        )


        # ====================================================
        # INSERIMENTO NELLA FINESTRA
        # ====================================================

        outer_layout.addWidget(
            logo_frame
        )

        outer_layout.addWidget(
            self.container
        )


        # ====================================================
        # COLONNA SINISTRA
        # ====================================================

        left_layout = QVBoxLayout()


        # Layer particelle (non modificabile)

        lbl_parcels = QLabel(
            "<b>Layer particelle</b>"
        )

        left_layout.addWidget(
            lbl_parcels
        )


        self.cmb_parcels = QComboBox()


        if (
            self.active_layer
            and
            self.active_layer in self.parcel_layers
        ):

            self.cmb_parcels.addItem(
                self.active_layer.name(),
                self.active_layer.id()
            )

        else:

            for lyr in self.parcel_layers:

                self.cmb_parcels.addItem(
                    lyr.name(),
                    lyr.id()
                )


        self.cmb_parcels.setEnabled(
            False
        )


        left_layout.addWidget(
            self.cmb_parcels
        )


        # Ricerca layer poligonali

        self.search_box = QLineEdit()

        self.search_box.setPlaceholderText(
            "Cerca layer poligonali..."
        )

        self.search_box.textChanged.connect(
            self.filter_tree
        )

        left_layout.addWidget(
            self.search_box
        )


        # Espandi / Collassa

        btn_expand = QPushButton(
            "Espandi tutto"
        )

        btn_collapse = QPushButton(
            "Collassa tutto"
        )


        btn_expand.clicked.connect(
            lambda: self.tree.expandAll()
        )

        btn_collapse.clicked.connect(
            lambda: self.tree.collapseAll()
        )


        expand_layout = QHBoxLayout()

        expand_layout.addWidget(
            btn_expand
        )

        expand_layout.addWidget(
            btn_collapse
        )

        left_layout.addLayout(
            expand_layout
        )


        # Albero layer

        self.tree = QTreeWidget()

        self.tree.setHeaderHidden(
            True
        )

        left_layout.addWidget(
            self.tree
        )


        self.populate_tree()


        # Pulisci selezione

        btn_clear = QPushButton(
            "Pulisci selezione"
        )

        btn_clear.clicked.connect(
            self.clear_selection
        )

        left_layout.addWidget(
            btn_clear
        )


        # OK / Annulla

        btn_layout = QHBoxLayout()

        btn_ok = QPushButton(
            "OK"
        )

        btn_cancel = QPushButton(
            "Annulla"
        )


        btn_ok.clicked.connect(
            self.accept
        )

        btn_cancel.clicked.connect(
            self.reject
        )


        btn_layout.addWidget(
            btn_ok
        )

        btn_layout.addWidget(
            btn_cancel
        )

        left_layout.addLayout(
            btn_layout
        )


        main_layout.addLayout(
            left_layout,
            2
        )


        # ====================================================
        # COLONNA DESTRA
        # ====================================================

        right_layout = QVBoxLayout()

        main_layout.addLayout(
            right_layout,
            1
        )


        lbl_sel = QLabel(
            "<b>Layer poligonali selezionati</b>"
        )

        right_layout.addWidget(
            lbl_sel
        )


        self.preview_panel = QTextEdit()

        self.preview_panel.setReadOnly(
            True
        )

        self.preview_panel.setPlaceholderText(
            "Nessun layer selezionato"
        )

        right_layout.addWidget(
            self.preview_panel
        )


        lbl_info = QLabel(
            "<b>Informazioni layer corrente</b>"
        )

        right_layout.addWidget(
            lbl_info
        )


        self.info_panel = QTextEdit()

        self.info_panel.setReadOnly(
            True
        )

        self.info_panel.setPlaceholderText(
            "Seleziona un layer per vedere le informazioni"
        )

        right_layout.addWidget(
            self.info_panel
        )


        self.tree.itemChanged.connect(
            self.update_selected_preview
        )

        self.tree.itemClicked.connect(
            self.update_layer_info
        )


        self.update_selected_preview()


    # ========================================================
    # POPOLAMENTO ALBERO
    # ========================================================

    def populate_tree(self):

        self.tree.clear()

        saved = self.settings.value(
            "selected_layers",
            [],
            type=list
        )


        root = QgsProject.instance().layerTreeRoot()


        def add_group(
            group,
            parent_item
        ):

            group_item = QTreeWidgetItem(
                parent_item,
                [group.name()]
            )


            group_item.setFlags(
                group_item.flags()
                &
                ~Qt.ItemIsSelectable
            )


            group_item.setFlags(
                group_item.flags()
                &
                ~Qt.ItemIsUserCheckable
            )


            group_item.setIcon(
                0,
                QIcon(
                    ":/images/themes/default/mIconFolder.svg"
                )
            )


            for child in group.children():

                if child.nodeType() == 0:

                    add_group(
                        child,
                        group_item
                    )


                elif child.nodeType() == 1:

                    layer = child.layer()


                    if layer in self.polygon_layers:

                        item = QTreeWidgetItem(
                            group_item,
                            [layer.name()]
                        )


                        item.setFlags(
                            item.flags()
                            |
                            Qt.ItemIsUserCheckable
                        )


                        item.setCheckState(
                            0,
                            (
                                Qt.Checked
                                if layer.name() in saved
                                else Qt.Unchecked
                            )
                        )


                        item.setData(
                            0,
                            Qt.UserRole,
                            layer.id()
                        )


                        item.setIcon(
                            0,
                            QIcon(
                                ":/images/themes/default/"
                                "mIconPolygonLayer.svg"
                            )
                        )


        add_group(
            root,
            self.tree
        )


        self.tree.expandAll()


    # ========================================================
    # FILTRO
    # ========================================================

    def filter_tree(
        self,
        text
    ):

        text = text.lower()


        def filter_item(
            item
        ):

            visible = (
                text
                in
                item.text(0).lower()
            )


            for i in range(
                item.childCount()
            ):

                child_visible = filter_item(
                    item.child(i)
                )

                visible = (
                    visible
                    or
                    child_visible
                )


            item.setHidden(
                not visible
            )


            return visible


        for i in range(
            self.tree.topLevelItemCount()
        ):

            filter_item(
                self.tree.topLevelItem(i)
            )


    # ========================================================
    # PULISCI SELEZIONE
    # ========================================================

    def clear_selection(
        self
    ):

        def clear_item(
            item
        ):

            if item.flags() & Qt.ItemIsUserCheckable:

                item.setCheckState(
                    0,
                    Qt.Unchecked
                )


            for i in range(
                item.childCount()
            ):

                clear_item(
                    item.child(i)
                )


        for i in range(
            self.tree.topLevelItemCount()
        ):

            clear_item(
                self.tree.topLevelItem(i)
            )


        self.update_selected_preview()


    # ========================================================
    # PREVIEW
    # ========================================================

    def update_selected_preview(
        self
    ):

        selected = []


        def scan(
            item
        ):

            if item.flags() & Qt.ItemIsUserCheckable:

                if item.checkState(0) == Qt.Checked:

                    selected.append(
                        item.text(0)
                    )


            for i in range(
                item.childCount()
            ):

                scan(
                    item.child(i)
                )


        for i in range(
            self.tree.topLevelItemCount()
        ):

            scan(
                self.tree.topLevelItem(i)
            )


        self.preview_panel.setText(
            "\n".join(selected)
            if selected
            else
            "Nessun layer selezionato"
        )


    # ========================================================
    # INFO LAYER
    # ========================================================

    def update_layer_info(
        self,
        item,
        column
    ):

        if item.data(
            0,
            Qt.UserRole
        ) is None:

            self.info_panel.setText(
                "Nessun layer selezionato"
            )

            return


        layer_id = item.data(
            0,
            Qt.UserRole
        )


        layer = QgsProject.instance().mapLayer(
            layer_id
        )


        if not layer:

            self.info_panel.setText(
                "Layer non trovato"
            )

            return


        name = layer.name()

        crs = layer.crs().authid()

        provider = layer.dataProvider().name()

        feature_count = layer.featureCount()


        warnings = []


        if feature_count == 0:

            warnings.append(
                "⚠️ Layer vuoto"
            )


        if layer.crs() != QgsProject.instance().crs():

            warnings.append(
                "⚠️ CRS diverso dal progetto"
            )


        # ----------------------------------------------------
        # Compatibilità struttura CDU
        # ----------------------------------------------------

        required_cdu = [
            "TEMA",
            "ZONA",
            "DETTAGLIO",
            "NORME"
        ]


        field_names = [
            field.name().upper()
            for field in layer.fields()
        ]


        missing_cdu = [
            field_name
            for field_name in required_cdu
            if field_name not in field_names
        ]


        # ----------------------------------------------------
        # Stato struttura
        # ----------------------------------------------------

        if layer.geometryType() != 2:

            cdu_status = (
                "⚠️ STRUTTURA NON COMPATIBILE CON IL CDU\n\n"
                "Il layer non è poligonale."
            )


        elif missing_cdu:

            cdu_status = (
                "⚠️ STRUTTURA NON COMPATIBILE CON IL CDU\n\n"
                "Campi mancanti:\n"
                +
                "\n".join(
                    f"• {field_name}"
                    for field_name in missing_cdu
                )
            )


        else:

            cdu_status = (
                "✓ STRUTTURA COMPATIBILE CON IL CDU"
            )


        # ----------------------------------------------------
        # Avvisi generali
        # ----------------------------------------------------

        warn_text = (
            "\n".join(warnings)
            if warnings
            else
            "Nessun problema rilevato"
        )


        # ----------------------------------------------------
        # Visualizzazione informazioni
        # ----------------------------------------------------

        self.info_panel.setText(
            f"Nome: {name}\n"
            f"CRS: {crs}\n"
            f"Feature: {feature_count}\n"
            f"Provider: {provider}\n\n"
            f"{warn_text}\n\n"
            f"{cdu_status}"
        )


    # ========================================================
    # LAYER SELEZIONATI
    # ========================================================

    def selected_layers(
        self
    ):

        selected_ids = []


        def scan(
            item
        ):

            if item.flags() & Qt.ItemIsUserCheckable:

                if item.checkState(0) == Qt.Checked:

                    selected_ids.append(
                        item.data(
                            0,
                            Qt.UserRole
                        )
                    )


            for i in range(
                item.childCount()
            ):

                scan(
                    item.child(i)
                )


        for i in range(
            self.tree.topLevelItemCount()
        ):

            scan(
                self.tree.topLevelItem(i)
            )


        self.settings.setValue(
            "selected_layers",
            [
                QgsProject.instance()
                .mapLayer(lid)
                .name()
                for lid in selected_ids
            ]
        )


        return [
            QgsProject.instance().mapLayer(lid)
            for lid in selected_ids
            if QgsProject.instance().mapLayer(lid)
        ]


    def selected_parcels_layer(
        self
    ):

        layer_id = (
            self.cmb_parcels.currentData()
        )

        return QgsProject.instance().mapLayer(
            layer_id
        )


# ============================================================
# VALIDAZIONI PARTICELLE
# ============================================================

def validate_parcels_structure(
    layer
):

    required = [
        "FID",
        "FOGLIO",
        "ALLEGATO",
        "MAPPALE",
        "SUPCALC"
    ]


    field_names = [
        f.name().upper()
        for f in layer.fields()
    ]


    for name in required:

        if name not in field_names:

            QMessageBox.critical(
                None,
                "Errore struttura particelle",
                f"Il layer particelle non contiene "
                f"il campo '{name}'."
            )

            return False


    return True


# ============================================================
# VALIDAZIONE POPOLAZIONE PARTICELLE
# ============================================================

def validate_parcels_population(
    layer
):

    required = [
        "FID",
        "FOGLIO",
        "ALLEGATO",
        "MAPPALE",
        "SUPCALC"
    ]


    for f in layer.getFeatures():

        for name in required:

            val = f[name]


            if val is None or val == "":

                QMessageBox.critical(
                    None,
                    "Errore dati particelle",
                    f"Il campo '{name}' contiene valori nulli."
                )

                return False


    return True


# ============================================================
# VALIDAZIONE STRUTTURA LAYER TEMATICI
# ============================================================

def validate_thematic_structure(
    layers
):

    required = [
        "TEMA",
        "ZONA",
        "DETTAGLIO",
        "NORME"
    ]


    for layer in layers:

        # ----------------------------------------------------
        # Controllo geometria
        # ----------------------------------------------------

        if (
            not hasattr(
                layer,
                "geometryType"
            )
            or
            layer.geometryType() != 2
        ):

            QMessageBox.critical(
                None,
                "Errore struttura layer tematico",
                f"Il layer:\n\n"
                f"'{layer.name()}'\n\n"
                f"non è un layer poligonale."
            )

            return False


        # ----------------------------------------------------
        # Controllo campi
        # ----------------------------------------------------

        field_names = [
            field.name().upper()
            for field in layer.fields()
        ]


        missing = [
            name
            for name in required
            if name not in field_names
        ]


        if missing:

            QMessageBox.critical(
                None,
                "Errore struttura layer tematico",
                f"Il layer:\n\n"
                f"'{layer.name()}'\n\n"
                f"non è compatibile con il CDU.\n\n"
                f"Campi mancanti:\n"
                +
                "\n".join(
                    f"• {name}"
                    for name in missing
                )
            )

            return False


    return True


# ============================================================
# INTERSEZIONI + PERCENTUALI + RECORDS
# ============================================================

def run_intersections(
    parcels_layer,
    selected_layers
):

    # ========================================================
    # NORMALIZZAZIONE TEMATICA
    #
    # I layer originali nella TOC NON vengono modificati.
    #
    # I layer già 2D vengono utilizzati direttamente.
    #
    # I layer con Z o M vengono convertiti temporaneamente
    # in geometrie 2D tramite native:dropmzvalues.
    # ========================================================

    normalized_layers = []


    for thematic_layer in selected_layers:

        wkb_type = thematic_layer.wkbType()


        has_z = QgsWkbTypes.hasZ(
            wkb_type
        )


        has_m = QgsWkbTypes.hasM(
            wkb_type
        )


        # ----------------------------------------------------
        # Layer già 2D
        # ----------------------------------------------------

        if not has_z and not has_m:

            normalized_layers.append(
                thematic_layer
            )

            continue


        # ----------------------------------------------------
        # Layer con Z/M
        #
        # Creazione di una copia temporanea 2D.
        # Il layer originale non viene modificato.
        # ----------------------------------------------------

        try:

            result = processing.run(
                "native:dropmzvalues",
                {
                    "INPUT": thematic_layer,
                    "DROP_M_VALUES": True,
                    "DROP_Z_VALUES": True,
                    "OUTPUT": "TEMPORARY_OUTPUT"
                }
            )


            normalized_layer = result["OUTPUT"]


            if not normalized_layer:

                raise Exception(
                    "L'algoritmo non ha restituito "
                    "un layer."
                )


            # ------------------------------------------------
            # Controllo di sicurezza
            #
            # Il risultato deve essere realmente 2D.
            # ------------------------------------------------

            normalized_wkb = (
                normalized_layer.wkbType()
            )


            if (
                QgsWkbTypes.hasZ(
                    normalized_wkb
                )
                or
                QgsWkbTypes.hasM(
                    normalized_wkb
                )
            ):

                raise Exception(
                    "Il layer risultante contiene ancora "
                    "Z o M."
                )


            normalized_layers.append(
                normalized_layer
            )


        except Exception as e:

            QMessageBox.critical(
                None,
                "Errore normalizzazione geometria",
                "Errore durante l'eliminazione "
                "dei valori Z/M dal layer:\n\n"
                +
                thematic_layer.name()
                +
                "\n\n"
                +
                str(e)
            )

            raise


    # ========================================================
    # MERGE LAYER TEMATICI NORMALIZZATI
    # ========================================================

    try:

        merged = processing.run(
            "native:mergevectorlayers",
            {
                "LAYERS": normalized_layers,
                "CRS": parcels_layer.crs().authid(),
                "OUTPUT": "TEMPORARY_OUTPUT"
            }
        )["OUTPUT"]


    except Exception as e:

        QMessageBox.critical(
            None,
            "Errore fusione layer tematici",
            "Errore nella fusione dei layer tematici:\n\n"
            +
            str(e)
        )

        raise


    # ========================================================
    # SALVA PARTICELLE SELEZIONATE
    # ========================================================

    parcels_sel = processing.run(
        "native:saveselectedfeatures",
        {
            "INPUT": parcels_layer,
            "OUTPUT": "TEMPORARY_OUTPUT"
        }
    )["OUTPUT"]


    # ========================================================
    # AGGIUNGE FK_SRC
    # ========================================================

    parcels_sel.startEditing()


    parcels_sel.dataProvider().addAttributes(
        [
            QgsField(
                "FK_SRC",
                QVariant.Int
            )
        ]
    )


    parcels_sel.updateFields()


    fid_field = next(
        (
            n
            for n in parcels_sel.fields().names()
            if n.lower() == "fid"
        ),
        None
    )


    for f in parcels_sel.getFeatures():

        f["FK_SRC"] = f[fid_field]

        parcels_sel.updateFeature(
            f
        )


    parcels_sel.commitChanges()


    # ========================================================
    # INTERSEZIONE UNICA
    # ========================================================

    inter = processing.run(
        "native:intersection",
        {
            "INPUT": parcels_sel,
            "OVERLAY": merged,
            "OUTPUT": "TEMPORARY_OUTPUT"
        }
    )["OUTPUT"]


    # ========================================================
    # CALCOLO PERCENTUALI
    # ========================================================

    inter.startEditing()


    supcalc_field = next(
        (
            n
            for n in inter.fields().names()
            if n.upper().startswith("SUPCALC")
        ),
        None
    )


    if (
        "PERCENT"
        not in
        [
            n.upper()
            for n in inter.fields().names()
        ]
    ):

        inter.dataProvider().addAttributes(
            [
                QgsField(
                    "PERCENT",
                    QVariant.Int
                )
            ]
        )


    if (
        "PERCENT_V"
        not in
        [
            n.upper()
            for n in inter.fields().names()
        ]
    ):

        inter.dataProvider().addAttributes(
            [
                QgsField(
                    "PERCENT_V",
                    QVariant.String
                )
            ]
        )


    inter.updateFields()


    for f in inter.getFeatures():

        area = f.geometry().area()


        try:

            sup = float(
                str(
                    f[supcalc_field]
                ).replace(
                    ",",
                    "."
                )
            )


            percent = (
                round(
                    area / sup * 100
                )
                if sup > 0
                else 0
            )


        except:

            percent = 0


        f["PERCENT"] = percent


        f["PERCENT_V"] = (
            "<1"
            if percent == 0
            else str(percent)
        )


        inter.updateFeature(
            f
        )


    inter.commitChanges()


    # ========================================================
    # ESTRAZIONE RECORDS
    # ========================================================

    records = []


    for f in inter.getFeatures():

        rec = {

            "foglio": f["FOGLIO"],

            "allegato": f["ALLEGATO"],

            "mappale": f["MAPPALE"],

            "tema": f["TEMA"],

            "zona": f["ZONA"],

            "dettaglio": f["DETTAGLIO"],

            "norme": f["NORME"],

            "percent_v": f["PERCENT_V"],

            "fk_cat": f["FK_SRC"]

        }


        records.append(
            rec
        )


    return records


# ============================================================
# ORDINAMENTO RECORDS
# ============================================================

def parse_mappale(
    value
):

    v = str(
        value
    ).strip()


    if v.isdigit():

        return (
            0,
            int(v),
            ""
        )


    match = re.match(
        r"^(\d+)([A-Za-z]+)$",
        v
    )


    if match:

        return (
            1,
            int(
                match.group(1)
            ),
            match.group(2)
        )


    if v.isalpha():

        return (
            2,
            0,
            v
        )


    return (
        3,
        0,
        v
    )


def sort_records(
    records
):

    return sorted(
        records,
        key=lambda r: (
            int(
                r["foglio"]
            ),
            parse_mappale(
                r["mappale"]
            ),
            r["tema"],
            r["zona"],
            r["percent_v"]
        )
    )


# ============================================================
# COMPILAZIONE CDU
# ============================================================

def compile_cdu(
    records,
    applica_filtro,
    output_file_path
):

    from docx import Document

    from docx.shared import Pt

    from docx.enum.table import (
        WD_CELL_VERTICAL_ALIGNMENT
    )

    from docx.enum.text import (
        WD_ALIGN_PARAGRAPH
    )

    from docx.oxml import parse_xml


    template_path = os.path.join(
        QgsProject.instance().homePath(),
        "media",
        "00_CDU_Schema.docx"
    )


    doc = Document(
        template_path
    )


    table = doc.tables[0]


    headers = [
        "Fg.",
        "All.",
        "Map.",
        "Tema",
        "Zona",
        "Dettaglio",
        "Norme Specifiche",
        "Q.tà* %"
    ]


    allineamenti = [

        WD_ALIGN_PARAGRAPH.CENTER,

        WD_ALIGN_PARAGRAPH.CENTER,

        WD_ALIGN_PARAGRAPH.CENTER,

        WD_ALIGN_PARAGRAPH.LEFT,

        WD_ALIGN_PARAGRAPH.CENTER,

        WD_ALIGN_PARAGRAPH.LEFT,

        WD_ALIGN_PARAGRAPH.LEFT,

        WD_ALIGN_PARAGRAPH.CENTER

    ]


    # ========================================================
    # INTESTAZIONE
    # ========================================================

    def aggiungi_intestazione():

        header_row = table.add_row()


        for i, h in enumerate(
            headers
        ):

            cell = header_row.cells[i]

            cell.text = h


            for p in cell.paragraphs:

                r = p.runs[0]

                r.bold = True

                r.font.name = "Calibri Light"

                r.font.size = Pt(10)

                p.alignment = allineamenti[i]


            # Rimuovi sfondo

            tcPr = (
                cell._element
                .get_or_add_tcPr()
            )


            for shd in tcPr.findall(
                ".//w:shd",
                {
                    "w":
                    "http://schemas.openxmlformats.org/"
                    "wordprocessingml/2006/main"
                }
            ):

                tcPr.remove(
                    shd
                )


            cell.vertical_alignment = (
                WD_CELL_VERTICAL_ALIGNMENT.CENTER
            )


    codici_ammessi = (
        "T01",
        "T02",
        "T03"
    )


    ultimo_gruppo = None


    # ========================================================
    # POPOLAMENTO TABELLA
    # ========================================================

    for r in records:

        if (
            applica_filtro
            and
            not str(
                r["tema"]
            ).startswith(
                codici_ammessi
            )
        ):

            continue


        foglio = int(
            r["foglio"]
        )


        allegato = r["allegato"]

        mappale = r["mappale"]


        gruppo = (
            foglio,
            mappale
        )


        # ----------------------------------------------------
        # NUOVO GRUPPO FOGLIO-MAPPALE
        # ----------------------------------------------------

        if ultimo_gruppo != gruppo:

            sep_row = table.add_row()


            first_cell = sep_row.cells[0]


            for c in sep_row.cells[1:]:

                first_cell.merge(
                    c
                )


            first_cell.text = (
                f"Foglio {foglio} – "
                f"Mappale {mappale}"
            )


            for p in first_cell.paragraphs:

                r0 = p.runs[0]

                r0.bold = True

                r0.font.name = "Calibri"

                r0.font.size = Pt(11)

                p.paragraph_format.space_before = Pt(6)


            # Rimuovi sfondo

            tcPr = (
                first_cell._element
                .get_or_add_tcPr()
            )


            for shd in tcPr.findall(
                ".//w:shd",
                {
                    "w":
                    "http://schemas.openxmlformats.org/"
                    "wordprocessingml/2006/main"
                }
            ):

                tcPr.remove(
                    shd
                )


            # Bordo inferiore

            borders = parse_xml(
                r'''
                <w:tcBorders xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
                    <w:top w:val="nil"/>
                    <w:bottom w:val="single" w:sz="6"/>
                    <w:left w:val="nil"/>
                    <w:right w:val="nil"/>
                </w:tcBorders>
                '''
            )


            tcPr.append(
                borders
            )


            first_cell.vertical_alignment = (
                WD_CELL_VERTICAL_ALIGNMENT.CENTER
            )


            # Intestazione dopo ogni gruppo

            aggiungi_intestazione()


            ultimo_gruppo = gruppo


        # ----------------------------------------------------
        # RIGA DATI
        # ----------------------------------------------------

        row = table.add_row()


        valori = [

            str(foglio),

            str(allegato),

            str(mappale),

            str(r["tema"]),

            str(r["zona"]),

            str(r["dettaglio"]),

            str(r["norme"]),

            str(r["percent_v"])

        ]


        for i, cell in enumerate(
            row.cells
        ):

            cell.text = valori[i]


            for p in cell.paragraphs:

                run = p.runs[0]

                run.font.name = "Calibri Light"

                run.font.size = Pt(10)

                p.alignment = allineamenti[i]


            cell.vertical_alignment = (
                WD_CELL_VERTICAL_ALIGNMENT.CENTER
            )


    # ========================================================
    # RIMOZIONE RIGA FANTASMA
    # ========================================================

    if (
        table.rows
        and
        all(
            cell.text.strip() == ""
            for cell in table.rows[0].cells
        )
    ):

        table._tbl.remove(
            table.rows[0]._tr
        )


    # ========================================================
    # SALVATAGGIO + APERTURA
    # ========================================================

    doc.save(
        output_file_path
    )


    if platform.system() == "Windows":

        os.startfile(
            output_file_path
        )

    elif platform.system() == "Darwin":

        os.system(
            f"open '{output_file_path}'"
        )

    elif platform.system() == "Linux":

        os.system(
            f"xdg-open '{output_file_path}'"
        )


# ============================================================
# WORKFLOW PRINCIPALE
# ============================================================

def run_workflow():

    try:

        # ----------------------------------------------------
        # Controllo dipendenza python-docx
        # ----------------------------------------------------

        if not ensure_python_docx():

            return


        # ----------------------------------------------------
        # Recupero layer dell'azione
        # ----------------------------------------------------

        layer_id = "[% @layer_id %]"


        active_layer = (
            QgsProject.instance()
            .mapLayer(layer_id)
        )


        if not active_layer:

            iface.messageBar().pushMessage(
                "Errore",
                "Impossibile identificare il layer dell'azione.",
                level=Qgis.Critical
            )

            return


        # ----------------------------------------------------
        # Recupero layer poligonali
        # ----------------------------------------------------

        polygon_layers = [

            lyr
            for lyr
            in
            QgsProject.instance()
            .mapLayers()
            .values()

            if (
                hasattr(
                    lyr,
                    "geometryType"
                )
                and
                lyr.geometryType() == 2
            )

        ]


        # ----------------------------------------------------
        # Dialogo completo
        # ----------------------------------------------------

        dialog = LayerSelectionDialog(
            polygon_layers=polygon_layers,
            parcel_layers=polygon_layers,
            output_tables=[],
            active_layer=active_layer
        )


        if dialog.exec_() != QDialog.Accepted:

            iface.messageBar().pushMessage(
                "Operazione annullata",
                "",
                level=Qgis.Info
            )

            return


        parcels_layer = (
            dialog.selected_parcels_layer()
        )


        selected_layers = (
            dialog.selected_layers()
        )


        # ----------------------------------------------------
        # Validazioni particelle
        # ----------------------------------------------------

        if not validate_parcels_structure(
            parcels_layer
        ):

            return


        if not validate_parcels_population(
            parcels_layer
        ):

            return


        # ----------------------------------------------------
        # Validazione layer tematici
        # ----------------------------------------------------

        if not validate_thematic_structure(
            selected_layers
        ):

            return


        # ----------------------------------------------------
        # Intersezioni → records
        # ----------------------------------------------------

        records = run_intersections(
            parcels_layer,
            selected_layers
        )


        if not records:

            iface.messageBar().pushMessage(
                "Nessun risultato",
                "Le particelle selezionate non intersecano "
                "i layer tematici.",
                level=Qgis.Warning
            )

            return


        # ====================================================
        # SCELTA ORDINARIO / COMPLETO
        # ====================================================

        msg = QMessageBox()


        msg.setWindowTitle(
            "Tipo di certificato"
        )


        msg.setText(
            "Desideri il certificato ORDINARIO "
            "(filtrato) o COMPLETO (tutti i dati)?"
        )


        btn_ordinario = msg.addButton(
            "ORDINARIO",
            QMessageBox.ActionRole
        )


        btn_completo = msg.addButton(
            "COMPLETO",
            QMessageBox.ActionRole
        )


        btn_annulla = msg.addButton(
            "ANNULLA",
            QMessageBox.RejectRole
        )


        msg.exec_()


        if msg.clickedButton() == btn_annulla:

            iface.messageBar().pushMessage(
                "Operazione annullata",
                "",
                level=Qgis.Info
            )

            return


        applica_filtro = (
            msg.clickedButton()
            ==
            btn_ordinario
        )


        # ====================================================
        # FINESTRA SALVATAGGIO
        # ====================================================

        output_file_path, _ = (
            QFileDialog.getSaveFileName(
                None,
                "Scegli dove salvare il CDU",
                "",
                "Documenti Word (*.docx)"
            )
        )


        if not output_file_path:

            iface.messageBar().pushMessage(
                "Operazione annullata",
                "",
                level=Qgis.Info
            )

            return


        # ====================================================
        # ORDINAMENTO RECORDS
        # ====================================================

        records_sorted = sort_records(
            records
        )


        # ====================================================
        # COMPILAZIONE CDU
        # ====================================================

        compile_cdu(
            records_sorted,
            applica_filtro,
            output_file_path
        )


        iface.messageBar().pushMessage(
            "CDU generato",
            "Documento salvato correttamente.",
            level=Qgis.Success
        )


    except Exception as e:

        iface.messageBar().pushMessage(
            "Errore critico",
            f"Si è verificato un errore: {str(e)}",
            level=Qgis.Critical
        )


        print(
            "❌ Errore:",
            e
        )


# ============================================================
# AVVIO SCRIPT
# ============================================================

run_workflow()
